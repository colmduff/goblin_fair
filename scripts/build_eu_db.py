"""Build a local SQLite database of the EU27 net-zero pathways, ready for FaIR.

Source: the ESABCC scenario database, release 2.0, REMIND 3.2 submission
(doi:10.5281/zenodo.8035686), saved in EU-data/ as
1687184510132-ESABCC_v2.0_ESABCC_REMIND_3.2.xlsx.
Its licence restricts redistribution, so the database is written to EU-data/,
which git ignores. EU-data/README.md documents the data and every choice below.

What it keeps: the EU27 rows of the three "NZero" scenarios. What it gives FaIR:
the four gas totals, as kt per year, interpolated to every year; and CH4 split
into biogenic and fossil streams from 1750, with PRIMAP-hist v2.6.1 history.

Run: make eu-db   (or: poetry run python scripts/build_eu_db.py)
"""

from __future__ import annotations

import sqlite3
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "EU-data" / "1687184510132-ESABCC_v2.0_ESABCC_REMIND_3.2.xlsx"
DATABASE = ROOT / "EU-data" / "eu27_net_zero.sqlite"

REGION = "EU27"
NZ_SCENARIOS = (
    "NZero_withICEPhOP",
    "NZero_bioLim12_withICEPhOP",
    "NZero_bioLim7p5_withICEPhOP",
)
# The report's names for the biomass cases (Byers et al. 2023, Table 2.2).
BIOMASS_LIMIT = {
    "NZero_withICEPhOP": "none (HiBio)",
    "NZero_bioLim12_withICEPhOP": "12 EJ/yr",
    "NZero_bioLim7p5_withICEPhOP": "7.5 EJ/yr",
}

# The only totals FaIR needs. CO2_FFI is the CO2 total minus AFOLU, so that
# CO2_FFI + CO2_AFOLU is exactly the database's own CO2 total.
CO2, AFOLU, CH4, N2O = (
    "Emissions|CO2",
    "Emissions|CO2|AFOLU",
    "Emissions|CH4",
    "Emissions|N2O",
)
TOTALS = (CO2, AFOLU, CH4, N2O)
FAIR_COLUMNS = ["CO2_FFI", "CO2_AFOLU", "CH4", "N2O"]
_TO_KT = {"Mt": 1000.0, "kt": 1.0}

# The sub-categories each gas is broken into (the only ones in the database).
PARTS = {
    "CH4": [
        "Emissions|CH4|AFOLU|Agriculture",
        "Emissions|CH4|AFOLU|Land",
        "Emissions|CH4|Energy|Supply",
        "Emissions|CH4|Waste",
    ],
    "N2O": [
        "Emissions|N2O|AFOLU|Agriculture",
        "Emissions|N2O|AFOLU|Land",
        "Emissions|N2O|Energy",
        "Emissions|N2O|Waste",
    ],
}


def read_iamc(path: Path = WORKBOOK, region: str = REGION) -> pd.DataFrame:
    """The `data` sheet, rows for one region only, years as int columns."""
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = book["data"].iter_rows(values_only=True)
    header = next(rows)
    keep = [row for row in rows if row[2] == region]
    book.close()
    frame = pd.DataFrame(keep, columns=header)
    return frame.rename(columns={c: int(c) for c in header[5:]})


def read_meta(path: Path = WORKBOOK) -> pd.DataFrame:
    """The `meta` sheet: one row per scenario."""
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(book["meta"].iter_rows(values_only=True))
    book.close()
    return pd.DataFrame(rows[1:], columns=rows[0])


def keep_scenarios(frame: pd.DataFrame, scenarios) -> pd.DataFrame:
    """Rows of these scenarios only, in the order given (not the file's order)."""
    order = {name: position for position, name in enumerate(scenarios)}
    kept = frame[frame["Scenario"].isin(order)]
    return kept.sort_values(
        "Scenario", key=lambda s: s.map(order), kind="stable"
    ).reset_index(drop=True)


def _years(frame: pd.DataFrame) -> list[int]:
    return [c for c in frame.columns if isinstance(c, int)]


def _series(frame: pd.DataFrame, scenario: str, variable: str) -> pd.Series:
    """One reported series in kt/yr, indexed by year, empty years dropped."""
    row = frame[(frame["Scenario"] == scenario) & (frame["Variable"] == variable)]
    if row.empty:
        raise ValueError(f"{scenario}: {variable} is missing")
    unit = str(row["Unit"].iloc[0])
    factor = _TO_KT.get(unit.split()[0])
    if factor is None:
        raise ValueError(f"{scenario}: {variable} has an unexpected unit {unit!r}")
    values = row[_years(frame)].iloc[0].astype(float).dropna()
    if values.empty:
        raise ValueError(f"{scenario}: {variable} has no values")
    return values * factor


def to_fair_input(frame: pd.DataFrame) -> pd.DataFrame:
    """The four FaIR inputs in kt/yr, for every year, one block per scenario.

    Reported years are kept exactly; the years between are linear
    interpolation. goblin_fair counts a missing year as zero emissions, so
    every year has to be filled.
    """
    blocks = []
    for scenario in frame["Scenario"].unique():
        totals = {v: _series(frame, scenario, v) for v in TOTALS}
        reported = totals[CO2].index
        years = np.arange(reported.min(), reported.max() + 1)

        def annual(s: pd.Series, years=years) -> np.ndarray:
            return np.interp(years, s.index.to_numpy(float), s.to_numpy())

        blocks.append(
            pd.DataFrame(
                {
                    "scenario": scenario,
                    "year": years,
                    "CO2_FFI": annual(totals[CO2]) - annual(totals[AFOLU]),
                    "CO2_AFOLU": annual(totals[AFOLU]),
                    "CH4": annual(totals[CH4]),
                    "N2O": annual(totals[N2O]),
                }
            )
        )
    return pd.concat(blocks, ignore_index=True)


def to_long(frame: pd.DataFrame) -> pd.DataFrame:
    """Every reported value as one row: scenario, variable, unit, year, value."""
    long = frame.melt(
        id_vars=["Scenario", "Variable", "Unit"],
        value_vars=_years(frame),
        var_name="year",
        value_name="value",
    ).dropna(subset=["value"])
    long.columns = ["scenario", "variable", "unit", "year", "value"]
    return long.sort_values(["scenario", "variable", "year"]).reset_index(drop=True)


def _classify(variable: str) -> tuple[str, str, str]:
    """(fair_input, part_of, note) for one variable name."""
    if variable == CO2:
        return "CO2_FFI + CO2_AFOLU", "", "CO2 total; CO2_FFI = this minus AFOLU"
    if variable == AFOLU:
        return "CO2_AFOLU", "", "net land CO2 (a sink, so negative)"
    if variable == CH4:
        return "CH4", "", "CH4 total"
    if variable == N2O:
        return "N2O", "", "N2O total"
    if "Bunkers" in variable:
        return "", "", "not used: international bunkers are outside the EU27 total"
    if variable == "Emissions|CO2|Share":
        return "", "", "not used: a share, not an emission"
    if variable.startswith(AFOLU + "|"):
        return "", "CO2_AFOLU", "sub-category, already inside the total"
    if variable.startswith(CO2 + "|"):
        return "", "CO2_FFI", "sub-category, already inside the total"
    for gas in ("CH4", "N2O"):
        if variable.startswith(f"Emissions|{gas}|"):
            return "", gas, "sub-category, already inside the total"
    if variable.startswith(("Emissions|F-Gases", "Emissions|Kyoto")):
        return "", "", "not used: given only in CO2-equivalent"
    if variable.startswith(
        ("Carbon Capture", "Carbon Removal", "Carbon Sequestration")
    ):
        return "", "", "not used: already netted into the CO2 totals"
    if variable.startswith("Gross Emissions"):
        return "", "", "not used: CO2 before removals"
    if variable.startswith("Diagnostics"):
        return "", "", "not used: harmonisation diagnostics"
    return "", "", "not used: energy-system or economic detail, not an emission"


def catalogue(frame: pd.DataFrame) -> pd.DataFrame:
    """Every variable, its unit and group, and what (if anything) it feeds."""
    names = frame[["Variable", "Unit"]].drop_duplicates("Variable")
    rows = []
    for variable, unit in names.itertuples(index=False):
        fair_input, part_of, note = _classify(variable)
        rows.append(
            {
                "variable": variable,
                "unit": unit,
                "group": variable.split("|")[0],
                "fair_input": fair_input,
                "part_of": part_of,
                "note": note,
            }
        )
    return pd.DataFrame(rows).sort_values("variable").reset_index(drop=True)


def breakdown_gaps(frame: pd.DataFrame) -> pd.DataFrame:
    """For CH4 and N2O: the total, the sum of its sub-categories, and the gap."""
    rows = []
    for scenario in frame["Scenario"].unique():
        one = frame[frame["Scenario"] == scenario].set_index("Variable")[_years(frame)]
        for gas, parts in PARTS.items():
            total = one.loc[f"Emissions|{gas}"].astype(float)
            present = [p for p in parts if p in one.index]
            summed = one.loc[present].astype(float).sum(min_count=1)
            for year in total.dropna().index:
                gap = total[year] - summed[year]
                rows.append(
                    {
                        "scenario": scenario,
                        "gas": gas,
                        "year": int(year),
                        "total": total[year],
                        "sum_of_parts": summed[year],
                        "gap": gap,
                        "gap_pct": 100 * gap / total[year],
                    }
                )
    return pd.DataFrame(rows)


def identical_to(frame: pd.DataFrame) -> dict[str, str]:
    """For each scenario, the first earlier scenario with exactly the same data."""
    seen: dict[str, pd.DataFrame] = {}
    out = {}
    for scenario in frame["Scenario"].unique():
        data = (
            frame[frame["Scenario"] == scenario]
            .set_index("Variable")[_years(frame)]
            .sort_index()
            .astype(float)
        )
        out[scenario] = next(
            (name for name, other in seen.items() if other.equals(data)), ""
        )
        seen[scenario] = data
    return out


def scenario_table(frame: pd.DataFrame, meta: pd.DataFrame) -> pd.DataFrame:
    same = identical_to(frame)
    columns = {
        "Category name": "category",
        "Vetting status": "vetting",
        "Feasibility Flag|Overall": "feasibility",
        "year of net-zero CO2 emissions (threshold=0 Gt CO2/yr)": "net_zero_co2_year",
        "year of net-zero GHGs full emissions (threshold=0 Gt CO2-equiv/yr)": (
            "net_zero_ghg_year"
        ),
        "Reference": "reference",
    }
    info = (
        meta.set_index("Scenario")
        .loc[list(same), list(columns)]
        .rename(columns=columns)
    )
    table = pd.DataFrame(
        {
            "scenario": list(same),
            "biomass_limit": [BIOMASS_LIMIT.get(s, "") for s in same],
            "identical_to": list(same.values()),
        }
    )
    return table.join(info, on="scenario").reset_index(drop=True)


# --- methane history (PRIMAP-hist) and biogenic / fossil streams -------------

# PRIMAP-hist v2.6.1 (Gütschow, Busch & Pflüger 2025), CC BY 4.0,
# doi:10.5281/zenodo.15016289. Main file: rounded to 3 significant digits.
PRIMAP_FILE = "Guetschow_et_al_2025-PRIMAP-hist_v2.6.1_final_13-Mar-2025.csv"
PRIMAP_URL = f"https://zenodo.org/records/15016289/files/{PRIMAP_FILE}?download=1"
PRIMAP_MD5 = "md5:09b9c61629f87e16012222e5b303bc36"
_P_SCENARIO = "scenario (PRIMAP-hist)"
_P_CATEGORY = "category (IPCC2006_PRIMAP)"
# Biogenic = agriculture (M.AG) + waste (4). Everything else in the national
# total excluding LULUCF (M.0.EL): energy (1), industry (2), other (5).
PRIMAP_BIOGENIC = ("M.AG", "4")
PRIMAP_FOSSIL = ("1", "2", "5")
_PRIMAP_OPTIONAL = ("5",)
# The main file is rounded, so the sectors match the total only to about 0.5 %.
_PRIMAP_TOLERANCE = 0.01
# ESABCC: biogenic = these parts; fossil = the CH4 total minus them, so the
# unitemised part (README §6) is counted as fossil and nothing is lost.
ESABCC_BIOGENIC = (
    "Emissions|CH4|AFOLU|Agriculture",
    "Emissions|CH4|AFOLU|Land",
    "Emissions|CH4|Waste",
)
JOIN_YEAR = 2005


def fetch_primap(folder: Path = WORKBOOK.parent) -> Path:
    """Download the PRIMAP-hist CSV once into EU-data/, checked against its md5."""
    import pooch

    return Path(
        pooch.retrieve(
            PRIMAP_URL, PRIMAP_MD5, fname=PRIMAP_FILE, path=folder, progressbar=False
        )
    )


def primap_ch4_streams(
    frame: pd.DataFrame, area: str = "EU27BX", scenario: str = "HISTCR"
) -> pd.DataFrame:
    """Biogenic and fossil CH4 for one area, kt/yr (1 Gg = 1 kt), by year."""
    rows = frame[
        (frame["area (ISO3)"] == area)
        & (frame["entity"] == "CH4")
        & (frame[_P_SCENARIO] == scenario)
    ].set_index(_P_CATEGORY)
    needed = (*PRIMAP_BIOGENIC, *PRIMAP_FOSSIL, "M.0.EL")
    # "Other" (5) is missing for some countries (e.g. Ireland): it is then zero.
    missing = [c for c in needed if c not in rows.index and c not in _PRIMAP_OPTIONAL]
    if missing:
        raise ValueError(f"PRIMAP-hist {area} CH4 has no categories {missing}")
    present = [c for c in needed if c in rows.index]
    units = set(rows.loc[present, "unit"])
    if units != {"CH4 * gigagram / yr"}:
        raise ValueError(f"unexpected PRIMAP-hist units {units}")
    years = [c for c in rows.columns if str(c).isdigit()]
    values = rows.loc[present, years].astype(float)
    fossil = [c for c in PRIMAP_FOSSIL if c in present]
    out = pd.DataFrame(
        {
            "biogenic": values.loc[list(PRIMAP_BIOGENIC)].sum(),
            "fossil": values.loc[fossil].sum(),
        }
    )
    out.index = pd.Index(out.index.astype(int), name="year")
    total = values.loc["M.0.EL"].to_numpy()
    off = np.abs(out.sum(axis=1).to_numpy() - total) / np.where(total, total, 1)
    if (off > _PRIMAP_TOLERANCE).any():
        raise ValueError(
            f"PRIMAP-hist sectors differ from M.0.EL by up to {off.max():.1%}"
        )
    return out


def esabcc_ch4_streams(frame: pd.DataFrame) -> pd.DataFrame:
    """Biogenic and fossil CH4 per scenario, kt/yr, every year (interpolated)."""
    blocks = []
    for scenario in frame["Scenario"].unique():
        total = _series(frame, scenario, CH4)
        years = np.arange(total.index.min(), total.index.max() + 1)

        def annual(s: pd.Series, years=years) -> np.ndarray:
            return np.interp(years, s.index.to_numpy(float), s.to_numpy())

        biogenic = sum(annual(_series(frame, scenario, v)) for v in ESABCC_BIOGENIC)
        blocks.append(
            pd.DataFrame(
                {
                    "scenario": scenario,
                    "year": years,
                    "biogenic": biogenic,
                    "fossil": annual(total) - biogenic,
                    "total": annual(total),
                }
            )
        )
    return pd.concat(blocks, ignore_index=True)


def join_history(
    history: pd.DataFrame, projection: pd.DataFrame, year: int = JOIN_YEAR
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """History before `year`, scaled per stream to meet the projection at `year`.

    Each stream's whole history is multiplied by one factor (projection / history
    in `year`), so the series joins without a jump and keeps its historical shape.
    """
    if year not in history.index or year not in projection.index:
        raise ValueError(f"the join year {year} must be in both history and projection")
    streams = ["biogenic", "fossil"]
    factors = projection.loc[year, streams] / history.loc[year, streams]
    before = history.loc[history.index < year, streams] * factors
    after = projection.loc[projection.index >= year, streams]
    joined = pd.concat([before, after])
    joined.index.name = "year"
    joined["total"] = joined["biogenic"] + joined["fossil"]
    joined["source"] = np.where(joined.index < year, "PRIMAP-hist", "ESABCC")
    table = pd.DataFrame(
        {
            "stream": streams,
            "join_year": year,
            "history_kt": history.loc[year, streams].to_numpy(),
            "projection_kt": projection.loc[year, streams].to_numpy(),
            "factor": factors.to_numpy(),
        }
    )
    return joined, table


def ch4_tables(frame: pd.DataFrame, primap: pd.DataFrame):
    """`ch4_streams` (1750-2100 per scenario) and `ch4_join` for the database."""
    history = primap_ch4_streams(primap)
    streams, joins = [], []
    for scenario, block in esabcc_ch4_streams(frame).groupby("scenario", sort=False):
        joined, factors = join_history(history, block.set_index("year"))
        streams.append(joined.reset_index().assign(scenario=scenario))
        joins.append(factors.assign(scenario=scenario))
    columns = ["scenario", "year", "biogenic", "fossil", "total", "source"]
    return (
        pd.concat(streams, ignore_index=True)[columns],
        pd.concat(joins, ignore_index=True),
    )


# --- Ireland: the Irish workbook and PRIMAP-hist IRL history -------------------

IRISH_WORKBOOK = WORKBOOK.parent / "ire_ch4-co2_temps_0.1.3.xlsx"
IRISH_JOIN_YEAR = 2020
# Irish workbook variable -> goblin_fair column (all in kt of the gas per year).
IRISH_VARIABLES = {
    "Emissions|CO2|MAGICC Fossil and Industrial": "CO2_FFI",
    "Emissions|CO2|MAGICC AFOLU": "CO2_AFOLU",
    "Emissions|CH4": "CH4",
    "Emissions|N2O": "N2O",
}
# MAGICC result blocks in sheet ire_temp_CH4_N2O_CO2 -> gas label.
IRISH_MAGICC_BLOCKS = {
    "Total Warming": "total",
    "Methane": "CH4",
    "N2O": "N2O",
    "CO2 Fossil": "CO2_FFI",
    "CO2 AFOLU": "CO2_AFOLU",
}


def _year_of(label) -> int:
    """'2020-01-01 00:00:00', a datetime, or '1/01/2020 0:00' -> 2020."""
    if hasattr(label, "year"):
        return int(label.year)
    text = str(label)
    for part in text.replace("-", " ").replace("/", " ").split():
        if len(part) == 4 and part.isdigit():
            return int(part)
    raise ValueError(f"no year in column label {label!r}")


def parse_irish_emissions(rows) -> pd.DataFrame:
    """Sheet ire_emissions (rows as tuples) -> scenario, year, four FaIR columns, kt.

    The sheet's region column says "World", but the numbers are Ireland's.
    """
    header = rows[0]
    years = [None if c is None else _year_of(c) for c in header[5:]]  # skip empty
    records = []
    for row in rows[1:]:
        if not row or row[4] not in IRISH_VARIABLES:
            continue
        if not str(row[3]).startswith("kt "):
            raise ValueError(f"unexpected Irish unit {row[3]!r} for {row[4]}")
        for year, value in zip(years, row[5:], strict=False):
            if year is not None and value is not None:
                records.append((row[2], year, IRISH_VARIABLES[row[4]], float(value)))
    long = pd.DataFrame(records, columns=["scenario", "year", "gas", "kt"])
    wide = long.pivot_table(index=["scenario", "year"], columns="gas", values="kt")
    return wide.reset_index()[
        ["scenario", "year", "CO2_FFI", "CO2_AFOLU", "CH4", "N2O"]
    ]


def parse_irish_magicc(rows) -> pd.DataFrame:
    """Sheet ire_temp_CH4_N2O_CO2 -> scenario, gas, year, warming_mK, background.

    The sheet holds blocks ("Total Warming", "Methane", ...), each a label row
    followed by MAGICC rows. A second, differently laid-out table follows the
    first set of blocks (its header starts "Year"); reading stops there.
    """
    header = rows[0]
    years = [None if c is None else _year_of(c) for c in header[7:]]  # skip empty
    records, gas, total_background = [], None, None
    for row in rows[1:]:
        if row[0] == "Year":
            break
        if row[0] is not None and all(v is None for v in row[1:]):
            gas = IRISH_MAGICC_BLOCKS.get(str(row[0]).strip())
            continue
        if gas is None or not str(row[1] or "").startswith("MAGICC"):
            continue
        if gas == "total":
            total_background = row[0]
        for year, value in zip(years, row[7:], strict=False):
            if year is not None and value is not None:
                records.append((row[4], gas, year, 1000 * float(value)))
    out = pd.DataFrame(records, columns=["scenario", "gas", "year", "warming_mK"])
    out["background"] = total_background  # per-gas blocks name the scenario instead
    return out


def read_irish(path: Path = IRISH_WORKBOOK) -> tuple[pd.DataFrame, pd.DataFrame]:
    """The Irish workbook: (emissions, MAGICC contributions)."""
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    emissions = parse_irish_emissions(
        list(book["ire_emissions"].iter_rows(values_only=True))
    )
    magicc = parse_irish_magicc(
        list(book["ire_temp_CH4_N2O_CO2"].iter_rows(values_only=True))
    )
    book.close()
    return emissions, magicc


def primap_series(
    frame: pd.DataFrame, area: str, entity: str, scenario: str = "HISTCR"
) -> pd.Series:
    """National total excluding LULUCF (M.0.EL) for one gas, kt/yr, by year."""
    rows = frame[
        (frame["area (ISO3)"] == area)
        & (frame["entity"] == entity)
        & (frame[_P_SCENARIO] == scenario)
        & (frame[_P_CATEGORY] == "M.0.EL")
    ]
    if len(rows) != 1:
        raise ValueError(f"PRIMAP-hist has {len(rows)} M.0.EL rows for {area} {entity}")
    if not str(rows["unit"].iloc[0]).endswith("gigagram / yr"):
        raise ValueError(f"unexpected PRIMAP-hist unit {rows['unit'].iloc[0]!r}")
    years = [c for c in rows.columns if str(c).isdigit()]
    out = rows[years].iloc[0].astype(float)
    out.index = pd.Index(out.index.astype(int), name="year")
    return out


def join_irish(
    scenario: pd.DataFrame, history: pd.DataFrame, year: int = IRISH_JOIN_YEAR
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """One Irish scenario from 1750: history before `year`, scaled to meet it.

    `scenario`: year and the four FaIR columns from `year` on. `history`:
    PRIMAP-hist by year for CO2_FFI, CH4 and N2O. CO2_AFOLU has no history in
    PRIMAP-hist's main file, so it is zero before `year` (stated in the README).
    """
    first = scenario.set_index("year").loc[year]
    factors, before = [], pd.DataFrame(index=history.index[history.index < year])
    for gas in ("CO2_FFI", "CH4", "N2O"):
        if year not in history.index:
            raise ValueError(f"the join year {year} must be in the history")
        factor = first[gas] / history.loc[year, gas]
        before[gas] = history.loc[history.index < year, gas] * factor
        factors.append(
            {"gas": gas, "join_year": year, "history_kt": history.loc[year, gas],
             "irish_kt": first[gas], "factor": factor}
        )  # fmt: skip
    before["CO2_AFOLU"] = 0.0
    before["source"] = "PRIMAP-hist"
    after = (
        scenario.set_index("year")
        .drop(columns="scenario")
        .assign(source="Irish workbook")
    )
    joined = pd.concat([before, after])[
        ["CO2_FFI", "CO2_AFOLU", "CH4", "N2O", "source"]
    ]
    joined.index.name = "year"
    return joined.reset_index(), pd.DataFrame(factors)


def irish_tables(emissions: pd.DataFrame, magicc: pd.DataFrame, primap: pd.DataFrame):
    """`ie_emissions` (1750-2100 per scenario), `ie_join` and `ie_magicc`."""
    history = pd.DataFrame(
        {gas: primap_series(primap, "IRL", entity) for gas, entity in
         [("CO2_FFI", "CO2"), ("CH4", "CH4"), ("N2O", "N2O")]}
    )  # fmt: skip
    joined, joins = [], []
    for name, block in emissions.groupby("scenario", sort=False):
        one, factors = join_irish(block, history)
        joined.append(one.assign(scenario=name))
        joins.append(factors.assign(scenario=name))
    columns = ["scenario", "year", "CO2_FFI", "CO2_AFOLU", "CH4", "N2O", "source"]
    return (
        pd.concat(joined, ignore_index=True)[columns],
        pd.concat(joins, ignore_index=True),
        magicc,
    )


def write_db(path: Path, tables: dict[str, pd.DataFrame]) -> None:
    with sqlite3.connect(path) as con:
        for name, table in tables.items():
            table.to_sql(name, con, if_exists="replace", index=False)


def fossil_ch4_correlation(frame: pd.DataFrame) -> float:
    """How closely the CH4 gap follows fossil primary energy, all scenarios.

    Uses the model years 2020-2050, every scenario in `frame`.
    """
    gaps = breakdown_gaps(frame)
    ch4 = gaps[(gaps["gas"] == "CH4") & gaps["year"].between(2020, 2050)]
    fossil = frame[frame["Variable"].isin(
        ["Primary Energy|Coal", "Primary Energy|Gas", "Primary Energy|Oil"]
    )]  # fmt: skip
    fossil = fossil.groupby("Scenario")[_years(frame)].sum(min_count=1)
    energy = [
        fossil.loc[s, y] for s, y in zip(ch4["scenario"], ch4["year"], strict=True)
    ]
    return float(np.corrcoef(ch4["gap"], energy)[0, 1])


def main() -> None:
    print(f"reading {WORKBOOK.name} (about 40 s) ...")
    everything = read_iamc(WORKBOOK)
    meta = read_meta(WORKBOOK)
    nz = keep_scenarios(everything, NZ_SCENARIOS)
    missing = set(NZ_SCENARIOS) - set(nz["Scenario"])
    if missing:
        raise ValueError(f"scenarios not in the workbook: {sorted(missing)}")

    print("reading PRIMAP-hist (downloaded once, md5 checked) ...")
    primap = pd.read_csv(fetch_primap())
    ch4_streams, ch4_join = ch4_tables(nz, primap)
    print("reading the Irish workbook ...")
    ie_emissions, ie_join, ie_magicc = irish_tables(*read_irish(), primap)

    tables = {
        "scenarios": scenario_table(nz, meta),
        "variables": catalogue(nz),
        "raw": to_long(nz),
        "fair_input": to_fair_input(nz),
        "breakdown_gaps": breakdown_gaps(nz),
        "ch4_streams": ch4_streams,
        "ch4_join": ch4_join,
        "ie_emissions": ie_emissions,
        "ie_join": ie_join,
        "ie_magicc": ie_magicc,
    }
    write_db(DATABASE, tables)

    print(f"wrote {DATABASE.relative_to(ROOT)}")
    for name, table in tables.items():
        print(f"  {name:15s} {len(table):7d} rows")
    print()
    print(tables["scenarios"].to_string(index=False))
    used = tables["variables"]
    print()
    print(f"variables: {len(used)}; feeding FaIR directly: "
          f"{(used['fair_input'] != '').sum()}; inside a total: "
          f"{(used['part_of'] != '').sum()}")  # fmt: skip
    print(used.groupby("note").size().to_string())
    gaps = tables["breakdown_gaps"]
    first = gaps[gaps["scenario"] == NZ_SCENARIOS[0]]
    print()
    print(first.pivot(index="year", columns="gas", values=["gap", "gap_pct"]).round(1))
    print()
    print("Irish history joined at 2020 (PRIMAP-hist IRL scaled by these factors):")
    print(
        ie_join[ie_join["scenario"] == ie_join["scenario"].iloc[0]]
        .round(3)
        .to_string(index=False)
    )
    print()
    print("CH4 history joined at 2005 (history scaled by these factors):")
    print(ch4_join[ch4_join["scenario"] == NZ_SCENARIOS[0]].to_string(index=False))
    print()
    print(
        "CH4 gap vs fossil primary energy, all "
        f"{everything['Scenario'].nunique()} scenarios, 2020-2050: "
        f"r = {fossil_ch4_correlation(everything):.2f}"
    )


if __name__ == "__main__":
    main()
